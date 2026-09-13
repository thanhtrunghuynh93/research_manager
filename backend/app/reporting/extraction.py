"""Turning an uploaded file into searchable text (REP-04).

One rule shapes the whole module: a file we could not read is marked unread, never indexed as
empty. Empty text and unreadable text lead to opposite conclusions about the week — one says there
was nothing to show, the other says we could not see it — and an assessment built on the wrong one
is unfair in a way nobody would notice.

So there are three outcomes rather than two. `ok` means we have the text. `unsupported` means this
kind of file has no text to extract and that is expected: a figure is evidence whether or not OCR
ever arrives. `failed` means we should have been able to read it and could not, which lowers the
evidence coverage of the assessment that relies on it (ASSESS-06).

Parsing runs on bytes we did not produce, so every parser is wrapped: a malformed file is a
`failed` extraction, never an exception that takes the job with it (requirements §11 "Security").
"""

from __future__ import annotations

import csv
import io
import json
import logging
from dataclasses import dataclass
from enum import StrEnum
from html.parser import HTMLParser

log = logging.getLogger(__name__)

# One artifact cannot be allowed to dominate the index or a prompt. A 200-page thesis is still
# 200 pages in the bucket; what is searchable is the first slice of it.
MAX_TEXT_CHARS = 200_000
MAX_PDF_PAGES = 200


class ExtractionState(StrEnum):
    PENDING = "pending"
    OK = "ok"
    FAILED = "failed"
    UNSUPPORTED = "unsupported"


@dataclass(frozen=True, slots=True)
class Extracted:
    state: ExtractionState
    text: str = ""
    truncated: bool = False
    # Why, in words a professor reading the coverage notes can use.
    note: str = ""


TEXT_EXTENSIONS = frozenset({"md", "txt", "tex", "bib", "json"})
HTML_EXTENSIONS = frozenset({"html", "htm"})
TABLE_EXTENSIONS = frozenset({"csv", "tsv"})
IMAGE_EXTENSIONS = frozenset({"png", "jpg", "jpeg", "gif", "webp", "svg"})
BINARY_EXTENSIONS = frozenset({"zip"})


def extract(filename: str, data: bytes) -> Extracted:
    extension = filename.rpartition(".")[2].lower()

    if extension in IMAGE_EXTENSIONS:
        # Not a failure: a figure is evidence, and OCR is explicitly later work (REP-04).
        return Extracted(
            state=ExtractionState.UNSUPPORTED,
            note="an image is stored as it is; no text was extracted and none was expected",
        )
    if extension in BINARY_EXTENSIONS:
        return Extracted(
            state=ExtractionState.UNSUPPORTED,
            note="an archive is stored as it is; its contents are not expanded",
        )

    try:
        if extension == "ipynb":
            return _cap(_notebook(data))
        if extension in TABLE_EXTENSIONS:
            return _cap(_table(data, delimiter="\t" if extension == "tsv" else ","))
        if extension == "pdf":
            return _cap(_pdf(data))
        if extension == "docx":
            return _cap(_docx(data))
        if extension == "pptx":
            return _cap(_pptx(data))
        if extension in HTML_EXTENSIONS:
            return _cap(_html(data))
        if extension in TEXT_EXTENSIONS or not extension:
            return _cap(_plain(data))
    except Exception as error:  # noqa: BLE001 - a malformed file is a state, not a crash
        log.warning("extraction failed for %s: %s", filename, type(error).__name__)
        return Extracted(
            state=ExtractionState.FAILED,
            note=f"the file could not be read ({type(error).__name__}); it is stored unchanged",
        )

    return Extracted(
        state=ExtractionState.UNSUPPORTED,
        note=f"no extractor for .{extension}; the file is stored unchanged",
    )


def _plain(data: bytes) -> Extracted:
    try:
        return Extracted(state=ExtractionState.OK, text=data.decode("utf-8"))
    except UnicodeDecodeError:
        # Mostly readable is worth having, and saying so is better than a silent substitution.
        return Extracted(
            state=ExtractionState.OK,
            text=data.decode("utf-8", errors="replace"),
            note="the file was not valid UTF-8; unreadable bytes were replaced",
        )


def _table(data: bytes, *, delimiter: str) -> Extracted:
    decoded = _plain(data)
    rows = list(csv.reader(io.StringIO(decoded.text), delimiter=delimiter))
    text = "\n".join(" | ".join(cell.strip() for cell in row) for row in rows if any(row))
    return Extracted(state=ExtractionState.OK, text=text, note=decoded.note)


def _notebook(data: bytes) -> Extracted:
    """Source and markdown only. Outputs are mostly base64 images, which are noise in an index."""
    document = json.loads(data.decode("utf-8"))
    parts: list[str] = []
    for cell in document.get("cells", []):
        source = cell.get("source", "")
        parts.append("".join(source) if isinstance(source, list) else str(source))
    return Extracted(state=ExtractionState.OK, text="\n\n".join(part for part in parts if part))


def _pdf(data: bytes) -> Extracted:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages[:MAX_PDF_PAGES]]
    text = "\n\n".join(page for page in pages if page.strip())
    if not text.strip():
        # A scanned PDF is a real case and is not a broken file; it simply has no text layer.
        return Extracted(
            state=ExtractionState.UNSUPPORTED,
            note="the PDF has no text layer, so nothing could be extracted without OCR",
        )
    note = (
        f"only the first {MAX_PDF_PAGES} pages were read"
        if len(reader.pages) > MAX_PDF_PAGES
        else ""
    )
    return Extracted(state=ExtractionState.OK, text=text, note=note)


def _docx(data: bytes) -> Extracted:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return Extracted(state=ExtractionState.OK, text="\n".join(parts))


def _pptx(data: bytes) -> Extracted:
    """Every text frame, slide by slide, plus the speaker notes.

    A deck is mostly headings and fragments, so slide boundaries are kept: "Results" under slide 4
    and "Results" under slide 9 are different claims, and a flat blob of phrases reads as one.
    Notes are included because that is where the argument usually is.
    """
    from pptx import Presentation

    parts: list[str] = []
    for number, slide in enumerate(Presentation(io.BytesIO(data)).slides, start=1):
        lines = [
            shape.text_frame.text.strip()
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        ]
        for table in (shape.table for shape in slide.shapes if shape.has_table):
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    lines.append(" | ".join(cells))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                lines.append(f"Notes: {notes.text.strip()}")
        if lines:
            parts.append(f"[slide {number}]\n" + "\n".join(lines))

    if not parts:
        # A deck of nothing but figures is evidence; it simply has no text to index.
        return Extracted(
            state=ExtractionState.UNSUPPORTED,
            note="the presentation has no text, so nothing could be extracted",
        )
    return Extracted(state=ExtractionState.OK, text="\n\n".join(parts))


class _TextOnly(HTMLParser):
    """Visible text only. Script and style bodies are code, and indexing them buries the prose."""

    SKIP = frozenset({"script", "style", "head", "title"})
    BREAK = frozenset({"p", "br", "div", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"})

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._muted = 0

    def handle_starttag(self, tag: str, attrs: object) -> None:
        if tag in self.SKIP:
            self._muted += 1

    def handle_endtag(self, tag: str) -> None:
        if tag in self.SKIP and self._muted:
            self._muted -= 1
        elif tag in self.BREAK:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._muted and data.strip():
            self.parts.append(data.strip())


def _html(data: bytes) -> Extracted:
    decoded = _plain(data)
    parser = _TextOnly()
    parser.feed(decoded.text)
    parser.close()
    text = " ".join(parser.parts).replace(" \n ", "\n").replace("\n ", "\n")
    lines = [line.strip() for line in text.split("\n")]
    return Extracted(
        state=ExtractionState.OK,
        text="\n".join(line for line in lines if line),
        note=decoded.note,
    )


def _cap(result: Extracted) -> Extracted:
    if len(result.text) <= MAX_TEXT_CHARS:
        return result
    note = "; ".join(
        part
        for part in (result.note, f"the text was truncated at {MAX_TEXT_CHARS:,} characters")
        if part
    )
    return Extracted(
        state=result.state, text=result.text[:MAX_TEXT_CHARS], truncated=True, note=note
    )
