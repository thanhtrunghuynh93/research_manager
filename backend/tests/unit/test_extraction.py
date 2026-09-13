"""Turning an uploaded file into searchable text (REP-04, REPO-08).

The rule that matters is the one about failure. A file we cannot read must be marked as unread,
not silently indexed as empty, because empty text and unreadable text lead to opposite conclusions
about the week: one says there was nothing, the other says we could not see.
"""

from __future__ import annotations

import io
import zipfile

import pytest

from app.reporting.extraction import ExtractionState, extract

pytestmark = pytest.mark.unit


def test_plain_text_and_markdown_come_through_as_written() -> None:
    result = extract("notes.md", b"# Week 3\n\nThe baseline reproduces.")

    assert result.state is ExtractionState.OK
    assert "The baseline reproduces" in result.text


def test_a_csv_keeps_its_rows_so_a_table_is_searchable() -> None:
    result = extract("results.csv", b"condition,loss\nbaseline,1.31\nablation,1.29\n")

    assert result.state is ExtractionState.OK
    assert "ablation" in result.text


def test_invalid_utf8_is_read_as_far_as_it_can_be_and_says_so() -> None:
    """A file with one bad byte is mostly readable, and mostly readable is worth having."""
    result = extract("notes.txt", "café".encode("latin-1") + b"\nrest of the note")

    assert result.state is ExtractionState.OK
    assert "rest of the note" in result.text
    assert result.note


def test_an_image_is_stored_without_text_rather_than_marked_broken() -> None:
    """REP-04: OCR comes later. A figure with no text is not an extraction failure."""
    result = extract("figure.png", b"\x89PNG\r\n\x1a\n" + b"\x00" * 32)

    assert result.state is ExtractionState.UNSUPPORTED
    assert result.text == ""
    assert "image" in result.note.lower()


def test_a_pdf_that_cannot_be_parsed_is_a_failure_not_an_empty_document() -> None:
    result = extract("paper.pdf", b"%PDF-1.7\nthis is not really a pdf")

    assert result.state is ExtractionState.FAILED
    assert result.text == ""
    assert result.note


def test_a_docx_that_cannot_be_parsed_is_a_failure_too() -> None:
    result = extract("draft.docx", b"PK\x03\x04 not really a docx")

    assert result.state is ExtractionState.FAILED


def test_a_real_docx_yields_its_paragraphs() -> None:
    document = _minimal_docx("The ablation did not help.")

    result = extract("draft.docx", document)

    assert result.state is ExtractionState.OK
    assert "ablation did not help" in result.text


def test_a_notebook_yields_its_source_and_not_its_base64_outputs() -> None:
    notebook = (
        b'{"cells": [{"cell_type": "code", "source": ["print(1)\\n"], '
        b'"outputs": [{"data": {"image/png": "iVBORw0KGgoAAAANS"}}]}]}'
    )

    result = extract("analysis.ipynb", notebook)

    assert result.state is ExtractionState.OK
    assert "print(1)" in result.text
    assert "iVBORw0" not in result.text


def test_extraction_is_capped_so_one_huge_file_cannot_fill_the_index() -> None:
    from app.reporting.extraction import MAX_TEXT_CHARS

    result = extract("big.txt", b"x" * (MAX_TEXT_CHARS + 5000))

    assert len(result.text) == MAX_TEXT_CHARS
    assert result.truncated
    assert result.state is ExtractionState.OK


def _minimal_docx(paragraph: str) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/'
            'package/2006/content-types"><Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" ContentType="application/vnd.'
            'openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>',
        )
        archive.writestr(
            "_rels/.rels",
            '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/'
            'package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.'
            'openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
            'Target="word/document.xml"/></Relationships>',
        )
        archive.writestr(
            "word/document.xml",
            '<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/'
            f'wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>{paragraph}</w:t></w:r>'
            "</w:p></w:body></w:document>",
        )
    return buffer.getvalue()


def _deck(title: str, bullets: str, notes: str = "") -> bytes:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = title
    slide.placeholders[1].text = bullets
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def test_a_pptx_yields_its_slides_and_speaker_notes() -> None:
    """A weekly package is often a deck, and the argument is usually in the notes."""
    deck = _deck("Retrieval baselines", "Recall@10 rose to 0.62", notes="Not conclusive yet.")

    result = extract("week3.pptx", deck)

    assert result.state is ExtractionState.OK
    assert "Retrieval baselines" in result.text
    assert "Recall@10 rose to 0.62" in result.text
    assert "Not conclusive yet" in result.text


def test_a_pptx_keeps_its_slide_boundaries() -> None:
    """ "Results" on slide 2 and "Results" on slide 9 are different claims, not one paragraph."""
    result = extract("week3.pptx", _deck("Results", "Seed spread dominates"))

    assert "[slide 1]" in result.text


def test_a_deck_of_only_figures_is_unsupported_rather_than_failed() -> None:
    from pptx import Presentation

    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])  # blank
    buffer = io.BytesIO()
    presentation.save(buffer)

    result = extract("figures.pptx", buffer.getvalue())

    assert result.state is ExtractionState.UNSUPPORTED
    assert "no text" in result.note


def test_a_pptx_that_cannot_be_parsed_is_a_failure_not_an_empty_deck() -> None:
    result = extract("week3.pptx", b"PK\x03\x04 not really a pptx")

    assert result.state is ExtractionState.FAILED


def test_html_yields_its_prose_without_its_markup() -> None:
    page = (
        b"<html><head><title>ignored</title><style>p{color:red}</style></head>"
        b"<body><h1>Weekly summary</h1><p>Curriculum ordering did not help.</p>"
        b"<script>var secret = 1;</script></body></html>"
    )

    result = extract("summary.html", page)

    assert result.state is ExtractionState.OK
    assert "Weekly summary" in result.text
    assert "Curriculum ordering did not help" in result.text
    # Script and style bodies are code. Indexed, they bury the prose they sit next to.
    assert "var secret" not in result.text
    assert "color:red" not in result.text
    assert "<p>" not in result.text


def test_html_entities_are_read_as_the_characters_they_name() -> None:
    result = extract("note.htm", b"<body><p>loss &lt; 1.3 &amp; falling</p></body>")

    assert result.state is ExtractionState.OK
    assert "loss < 1.3 & falling" in result.text
